#! /usr/bin/env python3

import contextlib
import io
import itertools
import json
import os
import shutil
import tempfile
import unittest
from types import SimpleNamespace
from unittest import mock

import dns.exception
import dns.resolver
import openpyxl
import requests

from spoofy import cli, master_table, report
from spoofy.dkim import DKIM
from spoofy.dmarc import DMARC, POLICIES, parse_tags, tree_walk_targets
from spoofy.master_table import SPREADSHEET, load_spreadsheet
from spoofy.master_table import TABLE as MASTER_TABLE
from spoofy.resolver import DNSResult, Resolver, get_resolver, txt_to_str
from spoofy.spf import SPF, parse_terms
from spoofy.spoofing import lookup


class FakeResolver:
    """Answers TXT lookups from a dict of record lists or DNSResults; anything else is NXDOMAIN."""

    nameservers = ("fake",)

    def __init__(self, txt=None, other=None):
        self.records = {
            name: v if isinstance(v, DNSResult) else DNSResult("ok", tuple(v))
            for name, v in (txt or {}).items()
        }
        self.other = other or {}

    def txt(self, name):
        return self.query(name, "TXT")

    def query(self, name, rdtype):
        if rdtype == "TXT" and name in self.records:
            return self.records[name]
        return self.other.get((name, rdtype), DNSResult("nxdomain"))


def run(spf="v=spf1 -all", dmarc=None, domain="example.com", records=None, exists=True):
    """process_domain against fake DNS: `spf` and `dmarc` are published for `domain`."""
    txt = dict(records or {})
    if spf:
        txt[domain] = [spf]
    if dmarc:
        txt[f"_dmarc.{domain}"] = [dmarc]
    other = {(domain, "A"): DNSResult("ok", ("192.0.2.1",))} if exists else {}
    return cli.process_domain(domain, resolver=FakeResolver(txt, other))


def code(*args, **kwargs):
    return run(*args, **kwargs)["SPOOFING_CODE"]


class TestDMARCParsing(unittest.TestCase):
    def test_sp_before_p_is_not_read_as_p(self):
        tags = parse_tags("v=DMARC1; sp=none; p=reject")
        self.assertEqual((tags["p"], tags["sp"]), ("reject", "none"))

    def test_np_is_not_read_as_p(self):
        self.assertEqual(parse_tags("v=DMARC1; np=reject; p=none")["p"], "none")

    def test_case_whitespace_and_duplicates(self):
        self.assertEqual(parse_tags("v=DMARC1; P = reject ; p=none;;")["p"], "reject")

    def test_tree_walk_targets(self):
        self.assertEqual(
            tree_walk_targets("a.b.example.com"),
            ["a.b.example.com", "b.example.com", "example.com", "com"],
        )
        long = "1.2.3.4.5.6.7.8.9.example.com"
        self.assertEqual(tree_walk_targets(long)[1], "5.6.7.8.9.example.com")


class TestDMARCDiscovery(unittest.TestCase):
    def dmarc(self, domain, records, other=None):
        return DMARC(domain, FakeResolver(records, other))

    def test_subdomain_own_record_wins(self):
        d = self.dmarc(
            "mail.example.com",
            {
                "_dmarc.mail.example.com": ["v=DMARC1; p=none"],
                "_dmarc.example.com": ["v=DMARC1; p=reject"],
            },
        )
        self.assertEqual(
            (d.record_domain, d.policy, d.inherited),
            ("mail.example.com", "none", False),
        )

    def test_inherited_record_and_domain_existence(self):
        records = {"_dmarc.example.com": ["v=DMARC1; p=reject; sp=none; np=quarantine"]}
        d = self.dmarc(
            "mail.example.com",
            records,
            {("mail.example.com", "A"): DNSResult("ok", ("192.0.2.1",))},
        )
        self.assertEqual(
            (d.inherited, d.sp, d.np, d.domain_exists),
            (True, "none", "quarantine", True),
        )
        self.assertFalse(self.dmarc("ghost.example.com", records).domain_exists)

    def test_public_suffix_domain_uses_its_own_record(self):
        d = self.dmarc(
            "gov.uk", {"_dmarc.gov.uk": ["v=DMARC1;p=reject;sp=none;np=reject"]}
        )
        self.assertEqual(d.policy, "reject")

    def test_multiple_records_are_discarded(self):
        d = self.dmarc(
            "mail.example.com",
            {
                "_dmarc.mail.example.com": ["v=DMARC1; p=none", "v=DMARC1; p=reject"],
                "_dmarc.example.com": ["v=DMARC1; p=quarantine"],
            },
        )
        self.assertEqual((d.record_domain, d.policy), ("example.com", "quarantine"))

    def test_organizational_domain_beats_intermediate_record(self):
        d = self.dmarc(
            "a.b.example.com",
            {
                "_dmarc.b.example.com": ["v=DMARC1; p=none"],
                "_dmarc.example.com": ["v=DMARC1; p=reject"],
            },
        )
        self.assertEqual((d.record_domain, d.policy), ("example.com", "reject"))

    def test_only_intermediate_record_is_used(self):
        d = self.dmarc(
            "a.b.example.com", {"_dmarc.b.example.com": ["v=DMARC1; p=quarantine"]}
        )
        self.assertEqual(d.record_domain, "b.example.com")

    def test_psd_n_marks_organizational_domain(self):
        d = self.dmarc(
            "a.b.example.com",
            {
                "_dmarc.b.example.com": ["v=DMARC1; p=quarantine; psd=n"],
                "_dmarc.example.com": ["v=DMARC1; p=reject"],
            },
        )
        self.assertEqual(d.record_domain, "b.example.com")

    def test_psd_y_falls_back_to_psd_record(self):
        records = {"_dmarc.gov": ["v=DMARC1; p=reject; np=reject; psd=y"]}
        d = self.dmarc("a.example.gov", records)
        self.assertEqual((d.record_domain, d.inherited), ("gov", True))
        records["_dmarc.example.gov"] = ["v=DMARC1; p=none"]
        self.assertEqual(
            self.dmarc("a.example.gov", records).record_domain, "example.gov"
        )

    def test_version_is_case_sensitive(self):
        d = self.dmarc("example.com", {"_dmarc.example.com": ["v=dmarc1; p=reject"]})
        self.assertIsNone(d.dmarc_record)
        self.assertTrue(any("malformed" in w for w in d.warnings))
        d = self.dmarc("example.com", {"_dmarc.example.com": ["V = DMARC1; p=reject"]})
        self.assertEqual(d.policy, "reject")

    def test_non_dmarc_txt_ignored(self):
        records = {
            "_dmarc.example.com": [
                "unrelated",
                "google-site-verification=DMARC1",
                "v=DMARC1; p=reject",
            ]
        }
        self.assertEqual(self.dmarc("example.com", records).policy, "reject")

    def test_invalid_policy_is_none_with_rua_else_ignored(self):
        records = {
            "_dmarc.example.com": ["v=DMARC1; p=rejected; rua=mailto:a@example.com"]
        }
        self.assertEqual(self.dmarc("example.com", records).policy, "none")
        records = {"_dmarc.example.com": ["v=DMARC1; p=rejected"]}
        self.assertIsNone(self.dmarc("example.com", records).policy)

    def test_invalid_aspf_falls_back_to_default(self):
        d = self.dmarc(
            "example.com", {"_dmarc.example.com": ["v=DMARC1; p=none; aspf=x"]}
        )
        self.assertEqual((d.aspf, len(d.warnings)), (None, 1))

    def test_lookup_error_is_not_no_dmarc(self):
        self.assertEqual(
            code(
                "v=spf1 ~all",
                records={"_dmarc.example.com": DNSResult("error", error="Timeout")},
            ),
            9,
        )


class TestResolver(unittest.TestCase):
    def test_errors_are_not_cached(self):
        r = Resolver(["192.0.2.1"])
        answers = [dns.exception.Timeout(), dns.resolver.NXDOMAIN()]

        def fake_resolve(*args, **kwargs):
            raise answers.pop(0)

        with mock.patch.object(r._resolver, "resolve", side_effect=fake_resolve):
            self.assertEqual(r.txt("example.com").status, "error")
            self.assertEqual(r.txt("example.com").status, "nxdomain")
            self.assertEqual(r.txt("example.com").status, "nxdomain")  # now cached

    def test_answers_are_parsed(self):
        r = Resolver(["192.0.2.1"])
        txt = SimpleNamespace(rrset=[SimpleNamespace(strings=[b"v=spf1 ", b"-all"])])
        ns = SimpleNamespace(
            rrset=[SimpleNamespace(to_text=lambda: "ns1.example.com.")]
        )
        with mock.patch.object(
            r._resolver, "resolve", side_effect=[SimpleNamespace(rrset=None), txt, ns]
        ):
            self.assertEqual(r.txt("a.example.com").status, "nodata")
            self.assertEqual(r.txt("b.example.com").records, ("v=spf1 -all",))
            self.assertEqual(
                r.query("example.com", "NS").records, ("ns1.example.com.",)
            )

    def test_resolvers_are_shared_per_server(self):
        self.assertIs(get_resolver(), get_resolver())
        self.assertEqual(get_resolver("192.0.2.53").nameservers, ["192.0.2.53"])
        self.assertIsNot(get_resolver("192.0.2.53"), get_resolver())

    def test_txt_strings_joined_without_spaces(self):
        rdata = SimpleNamespace(strings=[b"v=spf1 include:_spf.goo", b"gle.com ~all"])
        self.assertEqual(txt_to_str(rdata), "v=spf1 include:_spf.google.com ~all")


class TestSPF(unittest.TestCase):
    def spf(self, record, extra=None, other=None):
        return SPF(
            "example.com",
            FakeResolver({"example.com": [record], **(extra or {})}, other),
        )

    def test_all_mechanism(self):
        extra = {
            "mail-allow.example.net": ["v=spf1 -all"],
            "_spf.example.net": ["v=spf1 ~all"],
        }
        for record, want in [
            (
                "v=spf1 include:mail-allow.example.net ~all",
                "~all",
            ),  # hyphen is not a qualifier
            ("v=spf1 mx all", "+all"),
            ("v=spf1 mx -ALL", "-all"),
            ("v=spf1 -all ~all", "-all"),  # the first 'all' wins
            ("v=spf1 redirect=_spf.example.net", "~all"),
            (
                "v=spf1 -all redirect=_spf.example.net",
                "-all",
            ),  # redirect ignored with 'all'
        ]:
            self.assertEqual(self.spf(record, extra).all_mechanism, want, record)

    def test_lookup_counting(self):
        for record, want in [
            ("v=spf1 a mx -all", 2),
            ("v=spf1 -all a", 0),  # never evaluated after 'all'
            ("v=spf1 mx", 1),
            ("v=spf1 ~a ?mx -all", 2),
            ("v=spf1 a/24 mx/24 -all", 2),
            ("v=spf1 ptr:example.com -all", 1),
            ("v=spf1 a:x.com mx:y.com ptr exists:%{i}.z.com ip4:192.0.2.0/24 -all", 4),
            (
                "v=spf1 include:%{i}._spf.example.net -all exp=explain.example.net",
                1,
            ),  # macros aren't followed
        ]:
            self.assertEqual(self.spf(record).spf_dns_query_count, want, record)

    def test_warnings_are_not_repeated(self):
        self.assertEqual(len(self.spf("v=spf1 ptr ptr:example.com -all").warnings), 1)

    def test_nested_includes_counted(self):
        s = self.spf(
            "v=spf1 include:a.example.net -all",
            {
                "a.example.net": ["v=spf1 include:b.example.net mx ~all"],
                "b.example.net": ["v=spf1 a ~all"],
            },
        )
        self.assertEqual((s.spf_dns_query_count, s.errors), (4, []))

    def test_permerrors(self):
        includes = " ".join(f"include:i{n}.example.net" for n in range(11))
        extra = {f"i{n}.example.net": ["v=spf1 ip4:192.0.2.1 -all"] for n in range(11)}
        self.assertTrue(self.spf(f"v=spf1 {includes} -all", extra).too_many_dns_queries)
        for record in [
            "v=spf1 ip4 -all",
            "v=spf1 -all bogus",
        ]:  # syntax errors, even after 'all'
            self.assertTrue(self.spf(record).errors, record)
        loop = self.spf(
            "v=spf1 include:loop.example.net -all",
            {"loop.example.net": ["v=spf1 include:example.com ~all"]},
        )
        self.assertTrue(any("loop" in e for e in loop.errors))
        two = SPF(
            "example.com", FakeResolver({"example.com": ["v=spf1 -all", "v=spf1 ~all"]})
        )
        self.assertTrue(two.errors)
        extra = {
            "two.example.net": ["v=spf1 -all", "v=spf1 ~all"],
            "txt.example.net": ["verification=1"],
        }
        for record in [
            "v=spf1 include: -all",
            "v=spf1 include:two.example.net -all",
            "v=spf1 include:txt.example.net -all",
        ]:
            self.assertEqual(len(self.spf(record, extra).errors), 1, record)

    def test_include_chains_stop_at_max_depth(self):
        chain = {
            f"c{n}.example.net": [f"v=spf1 include:c{n + 1}.example.net -all"]
            for n in range(20)
        }
        s = self.spf("v=spf1 include:c0.example.net -all", chain)
        self.assertEqual((s.spf_dns_query_count, s.too_many_dns_queries), (11, True))

    def test_void_lookups_and_dangling_include(self):
        s = self.spf(
            "v=spf1 include:gone1.example.net include:gone2.example.net include:expired-vendor.net -all",
            other={("example.net", "NS"): DNSResult("ok", ("ns.example.net.",))},
        )
        self.assertEqual(
            (s.void_lookups, s.dangling_includes), (3, ["expired-vendor.net"])
        )
        self.assertTrue(s.errors)
        nodata = self.spf(
            "v=spf1 include:empty.example.net -all",
            {"empty.example.net": DNSResult("nodata")},
        )
        self.assertEqual((nodata.void_lookups, nodata.dangling_includes), (1, []))

    def test_failed_include_is_not_fatal_but_failed_record_is(self):
        r = FakeResolver({"example.com": ["v=spf1 include:flaky.example.net -all"]})
        r.records["flaky.example.net"] = DNSResult("error", error="Timeout")
        self.assertFalse(SPF("example.com", r).lookup_error)
        r.records["example.com"] = DNSResult("error", error="Timeout")
        self.assertTrue(SPF("example.com", r).lookup_error)

    def test_parse_terms(self):
        terms = parse_terms("v=spf1 -ip4:192.0.2.0/24 a/24 redirect=x.example")
        self.assertEqual(
            [(t.qualifier, t.name, t.value) for t in terms],
            [
                ("-", "ip4", "192.0.2.0/24"),
                ("+", "a", ""),
                (None, "redirect", "x.example"),
            ],
        )


SPF_RECORDS = {
    "-all": "v=spf1 -all",
    "~all": "v=spf1 ~all",
    "?all": "v=spf1 ?all",
    "+all": "v=spf1 +all",
    "noall": "v=spf1 ip4:192.0.2.1",
    "nospf": None,
}


def record_for(dmarc):
    """DMARC record for a table key: ('none', 'reject', None) -> 'v=DMARC1; p=none; sp=reject'."""
    if dmarc is None:
        return None
    return "v=DMARC1; " + "; ".join(
        f"{t}={v}" for t, v in zip(("p", "sp", "aspf"), dmarc) if v
    )


class TestMasterTable(unittest.TestCase):
    def test_module_is_exactly_what_regeneration_produces(self):
        with tempfile.TemporaryDirectory() as tmp:
            copy = shutil.copy(master_table.__file__, tmp)
            master_table.regenerate(copy, SPREADSHEET)
            with open(copy) as regenerated, open(master_table.__file__) as current:
                self.assertEqual(regenerated.read(), current.read())

    def test_duplicate_spreadsheet_rows_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook = openpyxl.Workbook()
            for row in (
                ["SPF", "DMARC", "Code"],
                ["-all", "p=none", 4],
                ["all-", "p=none", 4],
            ):
                workbook.active.append(row)
            workbook.save(os.path.join(tmp, "table.xlsx"))
            with self.assertRaises(ValueError):
                load_spreadsheet(os.path.join(tmp, "table.xlsx"))

    def test_every_row_is_reproduced(self):
        for (state, dmarc), expected in MASTER_TABLE.items():
            self.assertEqual(
                code(SPF_RECORDS[state], record_for(dmarc)), expected, (state, dmarc)
            )

    def test_every_record_resolves_to_a_tested_row(self):
        for state in {state for state, _ in MASTER_TABLE}:
            self.assertEqual(lookup(state), MASTER_TABLE[(state, None)])
            for p, sp, aspf in itertools.product(
                POLICIES, (None,) + POLICIES, (None, "r", "s")
            ):
                expected = MASTER_TABLE.get(
                    (state, (p, sp, aspf)), MASTER_TABLE.get((state, (p, p, aspf)))
                )
                self.assertIsNotNone(expected, (state, p, sp, aspf))
                self.assertEqual(lookup(state, p, sp, aspf), expected)

    def test_real_records_key_on_p_sp_aspf_as_written(self):
        record = "v=DMARC1; sp=none; rua=mailto:d@example.com; P=Reject; fo=1; adkim=s"
        self.assertEqual(
            code(dmarc=record), MASTER_TABLE[("-all", ("reject", "none", None))]
        )

    def test_pct_and_testing_mode(self):
        self.assertEqual(code(dmarc="v=DMARC1; p=quarantine; pct=50"), 3)
        # an invalid pct is ignored
        self.assertEqual(code(dmarc="v=DMARC1; p=quarantine; pct=abc"), 8)
        self.assertEqual(code(dmarc="v=DMARC1; p=reject; pct=50"), 8)
        self.assertEqual(code(dmarc="v=DMARC1; p=reject; t=y"), 8)
        self.assertEqual(code("v=spf1 ~all", "v=DMARC1; p=quarantine; t=y"), 0)
        self.assertEqual(code("v=spf1 ~all", "v=DMARC1; p=none; pct=0"), 0)

    def test_inherited_record_uses_the_parents_tested_subdomain_outcome(self):
        parent = {
            "example.com": ["v=spf1 -all"],
            "_dmarc.example.com": ["v=DMARC1; p=reject; sp=none; np=reject"],
        }
        # Table: -all | p=reject, sp=none -> 1 (subdomain spoofing possible)
        self.assertEqual(code(dmarc="v=DMARC1; p=reject; sp=none"), 1)
        result = run(None, domain="mail.example.com", records=parent)
        self.assertEqual(
            (result["DMARC_RECORD_DOMAIN"], result["SPOOFING_CODE"]), ("example.com", 0)
        )
        # np replaces sp for a nonexistent domain; the parent's SPF is the one that counts
        self.assertEqual(
            code(None, domain="ghost.example.com", records=parent, exists=False), 8
        )
        del parent["example.com"]  # No SPF | p=reject, sp=none -> 8
        self.assertEqual(code(None, domain="mail.example.com", records=parent), 8)
        partial = {"_dmarc.example.com": ["v=DMARC1; p=reject; sp=quarantine; pct=50"]}
        self.assertEqual(code(None, domain="mail.example.com", records=partial), 3)

    def test_spf_lookup_error_only_matters_without_enforcement(self):
        timeout = {"example.com": DNSResult("error", error="Timeout")}
        self.assertEqual(code(None, "v=DMARC1; p=reject", records=timeout), 8)
        self.assertEqual(code(None, "v=DMARC1; p=none", records=timeout), 9)

    def test_result_fields(self):
        result = run("v=spf1 ~all", "v=DMARC1; p=none; sp=reject")
        self.assertEqual(
            (result["SPOOFING_CODE"], result["SPOOFING_POSSIBLE"]), (2, True)
        )
        self.assertEqual(
            result["SPOOFING_TYPE"],
            "Organizational domain spoofing possible for example.com.",
        )


class TestCLI(unittest.TestCase):
    def test_worker_errors_do_not_escape(self):
        with mock.patch.object(cli, "process_domain", side_effect=RuntimeError("boom")):
            result = cli.safe_process_domain("example.com")
        self.assertEqual(
            (result["SPOOFING_CODE"], result["ERROR"]), (9, "RuntimeError: boom")
        )

    def test_result_is_json_serializable(self):
        bimi = {
            "default._bimi.example.com": [
                "unrelated",
                "v=BIMI1; l=https://example.com/logo.svg; a=",
            ]
        }
        result = run(dmarc="v=DMARC1; p=reject", records=bimi)
        json.dumps(result)
        self.assertEqual(
            (result["BIMI_LOCATION"], result["SPOOFING_CODE"]),
            ("https://example.com/logo.svg", 8),
        )

    def test_printer(self):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            report.printer(
                **run(
                    None,
                    domain="mail.example.com",
                    records={
                        "_dmarc.example.com": [
                            "v=DMARC1; p=reject; rua=mailto:d@example.com"
                        ]
                    },
                )
            )
        for line in [
            "No SPF record found.",
            "DMARC record inherited from example.com.",
            "Found DMARC policy: reject",
            "No DMARC pct found.",
            "Aggregate reports will be sent to: mailto:d@example.com",
            "Spoofing is not possible for mail.example.com.",
        ]:
            self.assertIn(line, out.getvalue())
        self.assertNotIn("non-existent subdomain policy", out.getvalue())
        # --dkim was not used, so the printer says nothing about DKIM
        self.assertNotIn("DKIM", out.getvalue())

    def test_dkim_line_gated_on_whether_it_was_checked(self):
        dns = FakeResolver(
            {
                "example.com": ["v=spf1 -all"],
                "_dmarc.example.com": ["v=DMARC1; p=reject"],
            }
        )
        # --dkim not used: the result records that, and the printer stays silent
        not_checked = cli.process_domain("example.com", resolver=dns)
        self.assertIs(not_checked["DKIM_CHECKED"], False)
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            report.printer(**not_checked)
        self.assertNotIn("DKIM", out.getvalue())
        # --dkim used but nothing found: the result records that, and the printer says so
        with mock.patch.object(cli, "DKIM") as fake_dkim:
            fake_dkim.return_value.dkim_record = None
            checked = cli.process_domain("example.com", enable_dkim=True, resolver=dns)
        self.assertIs(checked["DKIM_CHECKED"], True)
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            report.printer(**checked)
        self.assertIn("No known DKIM selectors enumerated", out.getvalue())

    def test_printer_full_and_error_results(self):
        result = run(
            "v=spf1 -all",
            records={
                "default._bimi.example.com": ["v=BIMI1; l=https://example.com/l.svg"]
            },
        )
        result.update(
            SPF_NUM_DNS_QUERIES=11,
            SPF_ERRORS=["11 DNS-querying terms (limit 10)"],
            SPF_DANGLING_INCLUDES=["gone.example"],
            DKIM="[*]    s1._domainkey.example.com -> k",
        )
        out = io.StringIO()
        with (
            contextlib.redirect_stdout(out),
            mock.patch.object(cli, "process_domain", side_effect=OSError("x")),
        ):
            report.printer(**result)
            report.printer(**cli.safe_process_domain("example.com"))
        for line in [
            "Too many SPF DNS query lookups 11.",
            "SPF permerror: 11 DNS-querying terms",
            "gone.example, which appears unregistered",
            "No DMARC record found.",
            "DKIM selectors:",
            "BIMI location: https://example.com/l.svg",
            "Error: OSError: x",
        ]:
            self.assertIn(line, out.getvalue())

    def test_dkim_keeps_latest_sighting_and_trims(self):
        dkim = DKIM.__new__(DKIM)
        dkim.domain = "example.com"
        formatted = dkim.format_dkim_records(
            [
                {"selector": "s1", "value": "A" * 200, "lastSeenAt": "2024-01-01"},
                {"selector": "s1", "value": "B" * 200, "lastSeenAt": "2025-01-01"},
                {"selector": "s1", "value": "C" * 200, "lastSeenAt": "2023-01-01"},
                {"selector": "s2", "domain": "x.example.com", "value": "short"},
                "junk",
            ]
        )
        self.assertEqual(
            formatted,
            f"[*]    s1._domainkey.example.com -> {'B' * 128}...(trimmed)\r\n"
            "[*]    s2._domainkey.x.example.com -> short",
        )
        self.assertIsNone(dkim.format_dkim_records({"error": "not a list"}))

    def test_dkim_api_failures_give_no_record(self):
        found = SimpleNamespace(
            status_code=200, json=lambda: [{"selector": "s1", "value": "k"}]
        )
        with mock.patch("requests.get", return_value=found):
            self.assertEqual(
                DKIM("example.com").dkim_record, "[*]    s1._domainkey.example.com -> k"
            )
        with mock.patch("requests.get", return_value=SimpleNamespace(status_code=429)):
            self.assertIsNone(DKIM("example.com").dkim_record)
        with mock.patch(
            "requests.get", side_effect=requests.exceptions.ConnectionError()
        ):
            self.assertIsNone(DKIM("example.com").dkim_record)

    def cli(self, *argv):
        dns = FakeResolver(
            {
                "example.com": ["v=spf1 -all"],
                "_dmarc.example.com": ["v=DMARC1; p=reject"],
            }
        )
        out = io.StringIO()
        with (
            mock.patch.object(cli, "get_resolver", return_value=dns),
            mock.patch("sys.argv", ["spoofy.py", *argv]),
            contextlib.redirect_stdout(out),
        ):
            cli.main()
        return out.getvalue()

    def test_cli_outputs(self):
        cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            os.chdir(tmp)
            try:
                with open("domains.txt", "w") as f:
                    f.write("example.com\nexample.org\n")
                results = json.loads(
                    self.cli("-iL", "domains.txt", "-o", "json", "-t", "2")
                )
                # results come back in input order
                domains = [r["DOMAIN"] for r in results]
                self.assertEqual(domains, ["example.com", "example.org"])
                self.assertIn(
                    "Spoofing is not possible for example.com.",
                    self.cli("-d", "example.com"),
                )
                self.cli("-iL", "domains.txt", "-o", "xls")
                self.cli("-d", "example.com", "-o", "xls")  # appends
                self.assertEqual(
                    openpyxl.load_workbook("output.xlsx").active.max_row, 4
                )
                open("empty.txt", "w").close()
                with (
                    self.assertRaises(SystemExit),
                    contextlib.redirect_stderr(io.StringIO()),
                ):
                    self.cli("-iL", "empty.txt")
            finally:
                os.chdir(cwd)

    def test_domain_list_normalization(self):
        path = os.path.join(os.path.dirname(__file__), ".test_domains.txt")
        with open(path, "w") as f:
            f.write(
                "Example.com\n\nhttps://example.com/path\nexample.org. # comment\n# only comment\n"
            )
        try:
            args = SimpleNamespace(d=None, iL=path)
            self.assertEqual(cli.read_domains(args), ["example.com", "example.org"])
        finally:
            os.remove(path)


if __name__ == "__main__":
    unittest.main()
